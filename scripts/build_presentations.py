"""Generate three editable AgentCacheX decks without loading original image assets."""

import argparse
from dataclasses import dataclass
from pathlib import Path
import tempfile
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
WIDTH, HEIGHT = 13.333333, 7.5
SOURCES = (
    "https://code.visualstudio.com/docs/agents/reference/workspace-context\n"
    "https://www.anthropic.com/engineering/writing-tools-for-agents\n"
    "https://www.anthropic.com/engineering/code-execution-with-mcp\n"
    "https://modelcontextprotocol.io/specification/2025-06-18/server/tools\n"
    "https://github.com/modelcontextprotocol/csharp-sdk\n"
    "https://git-scm.com/docs/git-grep\n"
    "https://redis.io/docs/latest/commands/set/"
)


@dataclass(frozen=True)
class Theme:
    filename: str
    label: str
    background: str
    panel: str
    foreground: str
    muted: str
    accent: str
    secondary: str
    border: str


THEMES = (
    Theme("AgentCacheX-Coding-Agent-Cache.pptx", "CODING AGENT",
          "F4F7FC", "FFFFFF", "13243D", "52647D", "0062CC", "007C69", "D9E3F0"),
    Theme("AgentCacheX-Keynote-Style.pptx", "KEYNOTE",
          "090909", "1B1B1B", "FAFAFA", "BBBBBB", "FFFFFF", "9ADACB", "383838"),
    Theme("AgentCacheX-Neon-Architecture.pptx", "NEON ARCHITECTURE",
          "100B24", "211737", "F7F1FF", "C4B7D7", "C486FF", "45E7CD", "51386E"),
)


def rgb(value):
    return RGBColor.from_string(value)


class Deck:
    def __init__(self, theme):
        self.theme = theme
        self.prs = Presentation()
        self.prs.slide_width = Inches(WIDTH)
        self.prs.slide_height = Inches(HEIGHT)
        self.prs.core_properties.title = "AgentCacheX - Shared cache. Smaller context."
        self.prs.core_properties.subject = "Proposed shared tool-result cache and token optimizer"
        self.prs.core_properties.author = "AgentCacheX"
        self.prs.core_properties.keywords = "MCP, Redis, progressive disclosure, language-agnostic"

    def text(self, slide, x, y, w, h, value, size=20, bold=False, color=None):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.text = value
        paragraphs = list(frame.paragraphs)
        for i, paragraph in enumerate(paragraphs):
            paragraph.font.name = "Segoe UI"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = rgb(color or self.theme.foreground)
            paragraph.line_spacing = 1.1
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(4 if i < len(paragraphs) - 1 else 0)
        return shape

    def panel(self, slide, x, y, w, h):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(self.theme.panel)
        shape.line.color.rgb = rgb(self.theme.border)
        shape.line.width = Pt(1)
        return shape

    def card(self, slide, x, y, w, h, title, body, title_size=24, body_size=20):
        self.panel(slide, x, y, w, h)
        self.text(slide, x + .24, y + .22, w - .48, .82,
                  title, title_size, True, self.theme.accent)
        self.text(slide, x + .24, y + 1.12, w - .48, h - 1.34,
                  body, body_size)

    def slide(self, section, title, notes):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(self.theme.background)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(WIDTH), Inches(.07)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(self.theme.accent)
        bar.line.fill.background()
        self.text(slide, .65, .35, 12, .3,
                  "AGENTCACHEX / " + section, 12, True, self.theme.accent)
        self.text(slide, .65, .9, 12.03, 1.08, title, 35, True)
        number = len(self.prs.slides)
        self.text(slide, .65, 7.1, 10.8, .25,
                  "PROPOSED MVP  |  Shared tool-result cache + token optimizer  |  " +
                  self.theme.label, 10, color=self.theme.muted)
        self.text(slide, 12.1, 7.05, .55, .35, f"{number:02d}", 14,
                  color=self.theme.muted)
        slide.notes_slide.notes_text_frame.text = (
            notes + "\n\nStatus: Proposed architecture, not an implemented service. "
            "Examples are illustrative, not measured product savings. "
            "See AgentCacheX-High-Level-Design.md for the full contract.\n\nSources:\n" + SOURCES
        )
        return slide

    def connector(self, slide, x1, y1, x2, y2, color=None):
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        line.line.color.rgb = rgb(color or self.theme.accent)
        line.line.width = Pt(2)
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        line.line._get_or_add_ln().append(end)

    def node(self, slide, x, y, w, title, body):
        self.panel(slide, x, y, w, 1.15)
        self.text(slide, x + .16, y + .14, w - .32, .38, title, 18, True,
                  self.theme.accent)
        self.text(slide, x + .16, y + .61, w - .32, .48, body, 15)

    def build(self):
        t = self.theme
        s = self.slide(
            "THE IDEA", "AgentCacheX",
            "Problem: multiple developers repeat read-only lookups against the same repository "
            "and send oversized results to their agents. Share version-qualified tool results "
            "and return only needed detail. Any programming language means a neutral provider-result "
            "boundary, not universal semantic understanding."
        )
        self.text(s, .65, 2.0, 12, 1.85, "Shared cache.\nSmaller context.", 48, True)
        self.text(s, .65, 4.08, 12, .65,
                  "One repository. Multiple developers. Any programming language.", 23,
                  color=t.muted)
        for x, title, body in (
            (.65, "SHARE RESULTS", "One MCP service + Redis"),
            (4.77, "SEND LESS", "Preview + selected details"),
            (8.89, "REUSE TOOLS", "No new code-analysis engine"),
        ):
            self.panel(s, x, 5.25, 3.8, 1.3)
            self.text(s, x + .22, 5.47, 3.36, .38, title, 19, True, t.accent)
            self.text(s, x + .22, 6.0, 3.36, .5, body, 16)

        s = self.slide(
            "THE PROBLEM", "Two costs. One focused problem.",
            "An exact cache alone can reduce repeated backend work while returning the same "
            "oversized payload. A token optimizer can help even on a cold cache miss. "
            "Existing providers may already cache and bound results; compare against those features."
        )
        self.card(s, .65, 2.25, 5.84, 3.0, "Duplicate lookups",
                  "We already searched this revision.\n"
                  "Another developer asks the same tool.", 27, 23)
        self.card(s, 6.84, 2.25, 5.84, 3.0, "Oversized context",
                  "The agent receives every match.\n"
                  "It only needs two source snippets.", 27, 23)
        self.text(s, .8, 5.85, 11.8, .7,
                  "Cache the work. Optimize what reaches the model.", 27, True, t.secondary)

        s = self.slide(
            "THE BOUNDARY", "Wrap existing tools. Do not rebuild them.",
            "Agents already have glob/file discovery, text search, reads, and sometimes semantic "
            "indexes, definitions, references, and call graphs. AgentCacheX reuses the chosen provider. "
            "The initial git-grep adapter produces lexical text matches, not proven references. "
            "Roslyn is C#-specific and is not an MVP dependency."
        )
        for x, title, body in (
            (.65, "Existing tools", "File search, text search,\nand source reads.\nKeep their native capabilities."),
            (4.77, "Our wrapper", "Shared result reuse.\nCompact previews.\nBatched selected details."),
            (8.89, "Not our engine", "No custom semantic analysis.\nNo language-server platform.\nNo generated evidence cards."),
        ):
            self.card(s, x, 2.25, 3.8, 3.7, title, body, 24, 20)
        self.text(s, .65, 6.25, 12, .6,
                  "Pilot: git grep at a verified commit. Text matches, not semantic references.",
                  18, color=t.muted)

        s = self.slide(
            "THE ARCHITECTURE", "One shared endpoint. Two useful paths.",
            "Developers use separate sessions/checkouts and explicitly configured MCP tools. "
            "The proposed ASP.NET Core/MCP host authenticates and authorizes each request and resolves "
            "the actual provider snapshot. A retained hit goes to the optimizer; a miss runs the "
            "existing provider, stores its full captured response, and then optimizes it too. "
            "Redis holds immutable payloads, pointers, expiry, and miss-coalescing leases. "
            "One service and one Redis instance demonstrate shared use, not high availability. "
            "MCP registration does not automatically intercept unrelated tools."
        )
        top = (
            (.65, "Developers A + B", "Separate agent sessions"),
            (3.75, "Shared MCP", "search_compact / details"),
            (6.85, "Access + snapshot", "Verify actual source"),
            (9.95, "Exact lookup", "Versioned result key"),
        )
        for x, title, body in top:
            self.node(s, x, 2.2, 2.73, title, body)
        for x in (3.43, 6.53, 9.63):
            self.connector(s, x, 2.78, x + .26, 2.78)
        self.node(s, .65, 4.65, 3.25, "Token optimizer", "Exact details to agents")
        self.node(s, 5.04, 4.65, 3.25, "Shared Redis", "Full results off-context")
        self.node(s, 9.43, 4.65, 3.25, "Existing provider", "Pinned-commit search / read")
        self.connector(s, 10.4, 3.35, 6.67, 4.6)
        self.text(s, 7.15, 3.83, 2.3, .4, "retained cache result", 15, color=t.muted)
        self.connector(s, 11.1, 3.35, 11.1, 4.6, t.secondary)
        self.text(s, 11.25, 3.87, 1.3, .4, "miss", 16, color=t.secondary)
        self.connector(s, 9.37, 5.23, 8.36, 5.23, t.secondary)
        self.text(s, 8.47, 4.86, .85, .35, "store", 14, color=t.muted)
        self.connector(s, 4.98, 5.23, 3.97, 5.23)
        self.text(s, .65, 6.25, 12, .55,
                  "Cache hits and misses both receive token-efficient output.", 22, True, t.secondary)

        s = self.slide(
            "TEAM REUSE", "Developer B benefits from Developer A.",
            "A hit requires the same repository, actual source/index snapshot, provider/tool version, "
            "canonical upstream arguments, effective access/policy scope, and result schema. "
            "Do not partition every key by username or preview budget. Include upstream limits and "
            "cursors that actually alter results. Changed commits create new keys, even for unrelated changes. "
            "Matching a paraphrased question is not a semantic-cache feature."
        )
        rows = (
            ("A: first lookup", "S1 + query Q + access scope P", "Run provider"),
            ("B: same lookup", "S1 + query Q + access scope P", "Shared hit"),
            ("A: code changed", "S2 + query Q + access scope P", "New key / miss"),
        )
        for index, (who, identity, result) in enumerate(rows):
            y = 2.2 + index * 1.18
            self.panel(s, .65, y, 12.03, .95)
            self.text(s, .9, y + .25, 2.7, .55, who, 20, True)
            self.text(s, 3.8, y + .25, 5.45, .55, identity, 20)
            self.text(s, 9.5, y + .25, 2.95, .55, result, 20, True, t.secondary)
        self.text(s, .65, 6.03, 12, .9,
                  "Different preview budgets can share a raw result.\n"
                  "Different permissions may require isolation.", 21, color=t.muted)

        s = self.slide(
            "TOKEN OPTIMIZATION", "Store everything captured. Send what is needed.",
            "ILLUSTRATIVE ARITHMETIC ONLY: 10,000 raw tool tokens versus 300 preview plus 1,200 "
            "selected-detail tokens is 1,500 emitted payload tokens, an 85% reduction in this example. "
            "This is not a measured result or total model-input/billed-cost reduction. Count all model "
            "turns, repeated history, prompts, tool definitions, output tokens, and provider pricing. "
            "If all details are needed, savings may disappear. Full captured response can be only one "
            "upstream page; preserve the provider's completeness and continuation metadata."
        )
        for x, value, label, detail in (
            (.65, "10,000", "STORED RAW", "Outside model context"),
            (4.77, "300", "PREVIEW", "Paths, IDs, short excerpts"),
            (8.89, "1,200", "DETAILS", "Only selected exact source"),
        ):
            self.panel(s, x, 2.2, 3.8, 2.5)
            self.text(s, x + .23, 2.43, 3.34, .35, label, 15, True, t.muted)
            self.text(s, x + .23, 2.96, 3.34, .9, value, 48, True, t.accent)
            self.text(s, x + .23, 4.01, 3.34, .6, detail, 18)
        self.text(s, .65, 5.05, 12, .65,
                  "300 + 1,200 = 1,500 emitted tool tokens", 30, True, t.secondary)
        self.text(s, .65, 6.0, 12, .75,
                  "Illustrative 85% tool-payload reduction.\n"
                  "Not measured. Not a total-task or billing guarantee.", 20, color=t.muted)

        s = self.slide(
            "TWO PROPOSED TOOLS", "Preview first. Fetch selected details together.",
            "Proposed signatures: search_compact(query, scope, response_budget, mode='preview', "
            "cursor=null); fetch_details(result_set_id, selections, response_budget, cursor=null). "
            "Both return source snapshot, expiry, coverage/continuation, and budget usage. "
            "Response budgets include metadata; exact token claims require the matching tokenizer, "
            "otherwise use an explicitly selected byte bound and labeled estimate. R17 and S1 are "
            "illustrative handles. For text not retained in the search result, perform and count an "
            "explicit source read pinned to the same revision, or report an unavailable snapshot. "
            "Never assume earlier excerpts survived compaction or delegation."
        )
        self.card(s, .65, 2.2, 5.84, 3.6, "search_compact",
                  "Query: VIP_DISCOUNT\nSnapshot: S1\n"
                  "Return: R17, item IDs, source previews\n"
                  "Show counts and remaining pages.", 26, 20)
        self.card(s, 6.84, 2.2, 5.84, 3.6, "fetch_details",
                  "Handle: R17\nSelections: r1 + r2 (one batch)\n"
                  "Return: exact text at S1\n"
                  "Report pending ranges when full.", 26, 20)
        self.text(s, .65, 6.15, 12, .6,
                  "R17 expired? Return RESULT_EXPIRED, then search again.", 23, True, t.secondary)

        s = self.slide(
            "CORRECTNESS", "A cache hit is useful only when it is safe.",
            "Verify actual searched revision, not model-supplied clean=true or commit hints. "
            "Dirty/unsaved or unknown workspace requests bypass shared reuse; if the provider cannot "
            "represent them, return UNSUPPORTED_SOURCE_STATE, not silently committed HEAD. "
            "Reauthorize every search/detail/page, including permission revocations. "
            "Snapshot mismatches, source unavailability, provider failures, and Redis failures are "
            "explicit. TTL is retention, not freshness. Preserve exact source text and meaningful "
            "provider fields. Complete mode must paginate, not silently top-k or claim grep is semantic."
        )
        for x, y, title, body in (
            (.65, 2.15, "Actual source snapshot",
             "Share verified committed results.\nDirty / unknown state: bypass or unsupported."),
            (6.84, 2.15, "Effective access scope",
             "Authorize every search, detail, and page.\nA result handle is not permission."),
            (.65, 4.25, "Honest completeness",
             "Preserve exact text and source locations.\nDisclose omitted items and partial pages."),
            (6.84, 4.25, "Explicit lifecycle",
             "TTL controls retention, not freshness.\nExpired handles require a new search."),
        ):
            self.panel(s, x, y, 5.84, 1.75)
            self.text(s, x + .23, y + .18, 5.38, .48, title, 23, True, t.accent)
            self.text(s, x + .23, y + .83, 5.38, .82, body, 18)

        s = self.slide(
            "HACKATHON SCOPE", "Build the narrow solution end to end.",
            "One repository and one existing provider. A .NET host is an implementation option, "
            "not a source-language restriction. Share a service, not a live database through OneDrive. "
            "Use bounded memory, TTL, entry size and response limits. Coalesce same-key misses with "
            "owner-checked expiring leases, bounded waits, and complete payload publication. "
            "This is not an exactly-once or high-availability claim. "
            "No owned Roslyn/LSP adapters, AST matching, embeddings, vector/graph database, evidence "
            "cards, dependency invalidation, or answer/patch/shell replay is needed."
        )
        self.card(s, .65, 2.2, 5.84, 4.0, "Build now",
                  "1. One pinned search/read provider\n"
                  "2. Shared MCP endpoint + Redis\n"
                  "3. Preview and detail response packing\n"
                  "4. Access, expiry, coalescing, metrics", 27, 21)
        self.card(s, 6.84, 2.2, 5.84, 4.0, "Do not build",
                  "Custom Roslyn / LSP adapters\n"
                  "Embeddings or vector / graph stores\n"
                  "Evidence cards or dependency graphs\n"
                  "Final-answer, edit, or shell replay", 27, 21)
        self.text(s, .65, 6.48, 12, .42,
                  "One service + one Redis instance prove team sharing, not high availability.",
                  18, color=t.muted)

        s = self.slide(
            "THE DEMO", "Prove reuse. Measure the whole task.",
            "Demonstrate two separate developer/agent sessions against one shared service. "
            "Cases: cold A, warm B with same and different preview budgets, changed snapshot, "
            "dirty/unknown state, unauthorized/revoked access, expiration, concurrent misses, "
            "complete-mode pagination, and oversized detail batches. Compare native bounded "
            "provider output, optimizer-only, and optimizer-plus-cache with the same source/model/task. "
            "Count backend searches AND pinned detail reads, all model input/output across all turns, "
            "latency, costs, and answer quality. Fast grep may not gain latency from Redis. "
            "Accept only correct access/source/completeness behavior, no material quality regression, "
            "and worthwhile measured benefit. No annual ROI or measured savings exists yet."
        )
        for x, title, body in (
            (.65, "Cold A -> warm B", "Same snapshot and access"),
            (4.77, "Change -> new key", "No stale current-state hit"),
            (8.89, "Expiry -> re-query", "No empty success fallback"),
        ):
            self.panel(s, x, 2.2, 3.8, 1.35)
            self.text(s, x + .22, 2.4, 3.36, .55, title, 21, True, t.accent)
            self.text(s, x + .22, 3.01, 3.36, .47, body, 17)
        self.card(s, .65, 3.95, 12.03, 2.65, "Decision gate",
                  "Compare native bounded tools, optimizer-only, and the shared cache.\n"
                  "Count all model turns, backend reads, latency, and answer quality.\n"
                  "Keep the wrapper only if its measured benefit is worthwhile.", 26, 21)
        return self.prs


def validate(path):
    # Inspect package names before loading: never read SVG/PNG or slide image assets.
    with ZipFile(path) as package:
        names = package.namelist()
        if any(name.lower().endswith((".svg", ".png")) or
               name.startswith("ppt/media/") for name in names):
            raise ValueError(f"Unexpected image asset in {path.name}")
    presentation = Presentation(path)
    if len(presentation.slides) != 10:
        raise ValueError(f"Expected ten slides in {path.name}")
    all_text = []
    for number, slide in enumerate(presentation.slides, 1):
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Missing speaker notes on slide {number}")
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or
                    shape.left + shape.width > presentation.slide_width + 10 or
                    shape.top + shape.height > presentation.slide_height + 10):
                raise ValueError(f"Shape out of bounds on slide {number}: {shape.name}")
            if shape.has_text_frame:
                all_text.append(shape.text)
    combined = "\n".join(all_text)
    for required in ("search_compact", "fetch_details", "Redis", "RESULT_EXPIRED",
                     "10,000", "1,200", "1,500", "Not measured"):
        if required not in combined:
            raise ValueError(f"Missing current-MVP content: {required}")
    for obsolete in ("257.4", "356,400", "8,000", "SQLite", "FTS5", "cache_lookup"):
        if obsolete in combined:
            raise ValueError(f"Obsolete architecture/claim in deck: {obsolete}")
    return len(presentation.slides)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    parser.add_argument("--check", action="store_true", help="Validate existing generated decks only")
    args = parser.parse_args()
    output = args.output_dir.resolve()
    if args.check:
        for theme in THEMES:
            count = validate(output / theme.filename)
            print(f"{theme.filename}: {count} slides, native shapes, current MVP")
        return
    output.mkdir(parents=True, exist_ok=True)
    # Validate every staged deck before replacing any existing presentation.
    with tempfile.TemporaryDirectory(prefix="agentcachex-slides-", dir=output) as temporary:
        staging = Path(temporary)
        for theme in THEMES:
            path = staging / theme.filename
            Deck(theme).build().save(path)
            validate(path)
        for theme in THEMES:
            (staging / theme.filename).replace(output / theme.filename)
            print(f"Updated {theme.filename}: 10 slides with speaker notes")


if __name__ == "__main__":
    main()
